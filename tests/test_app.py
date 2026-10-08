"""一時フォルダーだけを使う回帰テスト。実行: python test_app.py"""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import unittest
import tempfile
import hashlib
import time
from pathlib import Path
from test_support import make_test_symlink
from unittest.mock import patch
from PySide6.QtCore import Qt, QUrl
from PySide6.QtWidgets import QApplication, QMessageBox, QFileDialog
from app_main import App, HashScanWorker, MEDIA, QMediaPlayer, QTreeWidgetItemIterator
from safety import unique_planned, key, no_replace, identity, restore, validate_name

APP=QApplication.instance() or QApplication([])
def settle(window, seconds=5):
    end=time.monotonic()+seconds
    while window.busy() and time.monotonic()<end:
        APP.processEvents();time.sleep(.002)
    APP.processEvents()
    assert not window.busy(), '処理が終了しませんでした'

class SafetyTests(unittest.TestCase):
    def test_planned_collision(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'same.txt';p.write_text('x');reserved=set()
            a=unique_planned(p,reserved);reserved.add(key(a));b=unique_planned(p,reserved)
            self.assertNotEqual(a,b)
    def test_no_overwrite(self):
        with tempfile.TemporaryDirectory() as d:
            a,b=Path(d)/'a',Path(d)/'b';a.write_text('A');b.write_text('B')
            with self.assertRaises(FileExistsError):no_replace(a,b)
            self.assertEqual(a.read_text(),'A');self.assertEqual(b.read_text(),'B')
    def test_undo_conflict_preserves_retry(self):
        with tempfile.TemporaryDirectory() as d:
            a,b=Path(d)/'a',Path(d)/'b';a.write_text('A');entry=dict(current=str(b),original=str(a),identity=identity(a));no_replace(a,b);a.write_text('OTHER')
            remaining,errors=restore([entry]);self.assertEqual(len(remaining),1);self.assertTrue(errors);self.assertEqual(a.read_text(),'OTHER')
            a.unlink();remaining,errors=restore(remaining);self.assertFalse(remaining);self.assertFalse(errors);self.assertEqual(a.read_text(),'A')
    def test_modified_undo_refused(self):
        with tempfile.TemporaryDirectory() as d:
            a,b=Path(d)/'a',Path(d)/'b';a.write_text('A');entry=dict(current=str(b),original=str(a),identity=identity(a));no_replace(a,b);b.write_text('MODIFIED')
            remaining,errors=restore([entry]);self.assertTrue(errors);self.assertEqual(len(remaining),1);self.assertFalse(a.exists())
    def test_invalid_windows_names(self):
        for name in ['CON.txt','name?x.txt','../a','abc.','abc ','']:
            with self.assertRaises(ValueError):validate_name(name)

class HashTests(unittest.TestCase):
    def test_failure_and_chunk_progress(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'large.bin';p.write_bytes(b'x'*3*1024*1024);events=[];results=[];worker=HashScanWorker([p,Path(d)/'missing'])
            worker.progress.connect(events.append);worker.completed.connect(results.append);worker.run()
            rows=results[0]['records'];self.assertEqual(rows[0]['sha256'],hashlib.sha256(p.read_bytes()).hexdigest());self.assertEqual(rows[1]['status'],'失敗');self.assertTrue(rows[1]['error']);self.assertEqual(events[0]['done'],0)
    def test_cancel(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'large.bin';p.write_bytes(b'x'*4*1024*1024);result=[];worker=HashScanWorker([p]);worker.progress.connect(lambda _:worker.cancel(),Qt.DirectConnection);worker.completed.connect(result.append);worker.start();worker.wait(5000);APP.processEvents()
            self.assertTrue(result[0]['cancelled']);self.assertEqual(result[0]['records'][0]['status'],'未処理')

class GuiTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.env=patch.dict(os.environ,{'LOCALAPPDATA':self.tmp.name});self.env.start();self.w=App()
    def tearDown(self):
        settle(self.w);self.w.close()
        if MEDIA:self.w.player.setSource(QUrl())
        APP.processEvents();self.env.stop();self.tmp.cleanup()
    def test_preview_input_has_no_dialog(self):
        p=Path(self.tmp.name)/'abc.txt';p.write_text('ABC');self.w.files=[p]
        with patch.object(QMessageBox,'information') as dialog:
            self.w.rename_widgets['find'].setText('a');dialog.assert_not_called()
    def test_auto_avoid_and_undo(self):
        root=Path(self.tmp.name)/'files';root.mkdir()
        for name in ['a.txt','b.txt','same.txt']:(root/name).write_text(name)
        self.w.load(root);settle(self.w);self.w.rename_active='sequence';self.w.seq_pattern.setText('same');self.w.collision.setCurrentText('連番で自動回避');self.w.rename_preview()
        targets=[r['target'] for r in self.w.rename_rows];self.assertEqual(len(targets),len(set(targets)))
        with patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):self.w.rename_execute()
        settle(self.w);self.assertTrue(self.w.rename_undo);self.w.rename_undo_last();settle(self.w)
        self.assertEqual({p.name for p in root.iterdir()},{'a.txt','b.txt','same.txt'})
        for p in root.iterdir():self.assertEqual(p.read_text(),p.name)
    def test_hash_csv_uses_snapshot(self):
        root=Path(self.tmp.name)/'files';root.mkdir();p=root/'a.txt';p.write_text('X');self.w.load(root);settle(self.w);self.w.hash_scan();settle(self.w)
        p.unlink();out=Path(self.tmp.name)/'ledger.csv'
        with patch.object(QFileDialog,'getSaveFileName',return_value=(str(out),'CSV')):self.w.hash_export()
        settle(self.w);self.assertIn(hashlib.sha256(b'X').hexdigest(),out.read_text(encoding='utf-8-sig'));self.w.hash_clear();self.assertFalse(self.w.hash_records)
    def test_case_change(self):
        root=Path(self.tmp.name)/'files';root.mkdir();p=root/'abc.txt';p.write_text('X');self.w.load(root);settle(self.w);self.w.rename_active='case';self.w.case_mode.setCurrentText('大文字');self.w.rename_preview();self.assertEqual(self.w.rename_rows[0]['status'],'OK')
        with patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):self.w.rename_execute()
        settle(self.w);self.assertTrue((root/'ABC.txt').exists());self.w.rename_undo_last();settle(self.w);self.assertTrue(p.exists())
    def test_close_waits_for_hash(self):
        root=Path(self.tmp.name)/'files';root.mkdir();p=root/'large';p.write_bytes(b'x'*8*1024*1024);self.w.load(root);settle(self.w);self.w.hash_scan()
        with patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):self.w.close()
        settle(self.w);self.assertTrue(self.w.closing);self.assertFalse(self.w.hash_worker.isRunning())
    @unittest.skipUnless(MEDIA, 'メディア機能なし')
    def test_media_controls(self):
        self.w.media_controls.show();self.w.media_duration(120000);self.w.media_position(45000)
        self.assertEqual(self.w.seek_slider.maximum(),120000);self.assertEqual(self.w.seek_slider.value(),45000)
        self.w.volume_slider.setValue(25);self.assertAlmostEqual(self.w.audio.volume(),.25);self.w.mute_btn.click();self.assertTrue(self.w.audio.isMuted())
        self.w.media_playback_state(QMediaPlayer.PlayingState);self.assertTrue(self.w.pause_btn.isEnabled());self.assertFalse(self.w.play_btn.isEnabled())
        self.w.media_playback_state(QMediaPlayer.PausedState);self.assertTrue(self.w.play_btn.isEnabled());self.w.stop_btn.click();self.assertEqual(self.w.seek_slider.value(),0)
    @unittest.skipUnless(MEDIA, 'メディア機能なし')
    def test_media_visibility(self):
        import wave
        p=Path(self.tmp.name)/'sample.wav'
        with wave.open(str(p),'wb') as out:out.setnchannels(1);out.setsampwidth(2);out.setframerate(8000);out.writeframes(b'\0'*16000)
        self.w.preview(p);self.assertFalse(self.w.media_controls.isHidden());self.assertTrue(self.w.video.isHidden())
        text=Path(self.tmp.name)/'text.txt';text.write_text('Hello');self.w.preview(text);settle(self.w);self.assertTrue(self.w.media_controls.isHidden())
    def test_seek_click_drag_keys(self):
        from app_main import SeekSlider
        from PySide6.QtTest import QTest
        from PySide6.QtCore import QPoint
        slider=SeekSlider();slider.setRange(0,100000);slider.setSingleStep(1000);slider.resize(400,40);slider.show();APP.processEvents();requested=[];slider.seekRequested.connect(requested.append)
        QTest.mouseClick(slider,Qt.LeftButton,pos=QPoint(200,20));self.assertTrue(45000<slider.value()<55000)
        QTest.mousePress(slider,Qt.LeftButton,pos=QPoint(80,20));QTest.mouseMove(slider,QPoint(320,20));QTest.mouseRelease(slider,Qt.LeftButton,pos=QPoint(320,20));self.assertTrue(slider.value()>75000)
        QTest.keyClick(slider,Qt.Key_Home);self.assertEqual(slider.value(),0);QTest.keyClick(slider,Qt.Key_Right);self.assertEqual(slider.value(),1000);self.assertEqual(requested[-1],1000);slider.close()
    def test_media_edit_disabled(self):
        for extension in ('.mp4','.mp3'):
            p=Path(self.tmp.name)/('sample'+extension);p.write_bytes(b'')
            self.w.preview(p);self.assertFalse(self.w.edit_btn.isEnabled())
            with patch.object(self.w,'open_editor') as editor:self.w.edit_current();editor.assert_not_called()
        p=Path(self.tmp.name)/'sample.txt';p.write_text('text');self.w.preview(p);settle(self.w);self.assertTrue(self.w.edit_btn.isEnabled())
    def test_video_seek_commits_once(self):
        from unittest.mock import MagicMock
        player=self.w.player;fake=MagicMock();fake.duration.return_value=100000;fake.playbackState.return_value=QMediaPlayer.PlayingState;self.w.player=fake
        try:
            self.w.seek_slider.setRange(0,100000);self.w.seek_slider.setSliderDown(True);self.w.seek_slider.setValue(30000);self.w.media_seek(30000);self.w.seek_slider.setValue(70000);self.w.media_seek(70000)
            fake.setPosition.assert_not_called();fake.pause.assert_called_once();self.w.seek_slider.setSliderDown(False);fake.setPosition.assert_called_once_with(70000);fake.play.assert_called_once()
        finally:self.w.player=player
    def test_video_seek_keeps_pause_and_enables_after_duration(self):
        from unittest.mock import MagicMock
        player=self.w.player;fake=MagicMock();fake.duration.return_value=100000;fake.position.return_value=0;fake.playbackState.return_value=QMediaPlayer.PausedState;self.w.player=fake
        try:
            self.w.media_duration(100000);self.assertTrue(self.w.seek_slider.isEnabled());self.w.seek_slider.setSliderDown(True);self.w.seek_slider.setValue(50000);self.w.seek_slider.setSliderDown(False);fake.setPosition.assert_called_once_with(50000);fake.play.assert_not_called()
        finally:self.w.player=player
    def test_specification_and_icons(self):
        from app_main import resource_path,PDF_VIEW
        self.assertFalse(self.w.windowIcon().isNull());self.assertTrue(resource_path('assets/app.ico').exists());self.assertTrue(resource_path('FileWorkbench_Next_Specification_Ver2.7.0_Final.pdf').exists())
        if PDF_VIEW:
            self.w.spec_btn.click();APP.processEvents();self.assertIsNotNone(self.w.spec_dialog);self.assertTrue(self.w.spec_dialog.isVisible());self.w.spec_dialog.close()
    def fixture_tree(self):
        root=Path(self.tmp.name)/'source';root.mkdir();(root/'nested').mkdir();(root/'empty').mkdir()
        (root/'one.txt').write_text('ONE');(root/'nested'/'a&b.txt').write_text('TWO');return root
    def drop_on(self,target,paths):
        from PySide6.QtCore import QMimeData,QPoint,QPointF
        from PySide6.QtGui import QDragEnterEvent,QDropEvent
        mime=QMimeData();mime.setUrls([QUrl.fromLocalFile(str(path)) for path in paths]);enter=QDragEnterEvent(QPoint(10,10),Qt.CopyAction,mime,Qt.LeftButton,Qt.NoModifier);APP.sendEvent(target,enter);event=QDropEvent(QPointF(10,10),Qt.CopyAction,mime,Qt.LeftButton,Qt.NoModifier);APP.sendEvent(target,event);self.assertTrue(event.isAccepted());settle(self.w)
    def test_hash_unique_files_are_listed_by_default(self):
        root=self.fixture_tree();self.w.load(root);settle(self.w);self.w.hash_scan_btn.click();settle(self.w)
        self.assertEqual(self.w.hash_mode.currentText(),'全件台帳');self.assertEqual(self.w.hash_list.count(),2);self.assertEqual(len(self.w.hash_records),2)
        self.w.hash_mode.setCurrentText('重複のみ');settle(self.w);self.assertIn('重複ファイルはありません',self.w.hash_list.item(0).text())
        self.w.hash_mode.setCurrentText('全件台帳');settle(self.w);self.assertEqual(self.w.hash_list.count(),2)
    def test_hash_drop_folder_auto_scans(self):
        root=self.fixture_tree();self.w.show_page('hash');self.drop_on(self.w.hash_list.viewport(),[root]);self.assertEqual(len(self.w.hash_records),2);self.assertEqual(self.w.hash_list.count(),2);self.assertIs(self.w.stack.currentWidget(),self.w.pages['hash'])
    def test_hash_drop_multiple_files_auto_scans(self):
        root=self.fixture_tree();self.w.show_page('hash');self.drop_on(self.w.hash_list.viewport(),[root/'one.txt',root/'nested'/'a&b.txt',root/'one.txt']);self.assertEqual(len(self.w.hash_records),2);self.assertEqual(self.w.hash_list.count(),2)
    def test_main_drop_loads_tree(self):
        root=self.fixture_tree();self.drop_on(self.w,[root]);self.assertEqual(self.w.root,root);self.assertEqual(len(self.w.files),2);self.assertEqual(len(self.w.tree_records),5);self.assertIs(self.w.stack.currentWidget(),self.w.pages['tree'])
    def test_drop_on_path_field_loads_tree(self):
        root=self.fixture_tree();self.drop_on(self.w.tree_root_edit,[root]);self.assertEqual(len(self.w.files),2)
    def test_drop_compare_left_and_right(self):
        left=self.fixture_tree();right=Path(self.tmp.name)/'right';right.mkdir();(right/'one.txt').write_text('ONE');self.w.show_page('compare');self.drop_on(self.w.left.viewport(),[left]);self.drop_on(self.w.right.viewport(),[right]);self.w.compare();settle(self.w);self.assertTrue(any('完全一致' in self.w.left.item(i).text() for i in range(self.w.left.count())))
    def test_tree_outputs_all_formats(self):
        import csv
        from openpyxl import load_workbook
        root=self.fixture_tree();self.w.load(root);settle(self.w);self.w.tree_save_dir.setText(self.tmp.name);self.w.tree_numbering.setChecked(True)
        for format_name,extension in [('TXT','.txt'),('CSV','.csv'),('Excel','.xlsx'),('HTML','.html')]:
            self.w.tree_format.setCurrentText(format_name);self.w.tree_output_name.setText('structure')
            with patch.object(QMessageBox,'information'):self.w.tree_export_btn.click();settle(self.w)
            out=Path(self.tmp.name)/('structure'+extension);self.assertTrue(out.exists())
        with (Path(self.tmp.name)/'structure.csv').open(encoding='utf-8-sig') as f:rows=list(csv.reader(f));self.assertEqual(len(rows),6)
        wb=load_workbook(Path(self.tmp.name)/'structure.xlsx');self.assertEqual(wb.active.max_row,6);wb.close()
        html=(Path(self.tmp.name)/'structure.html').read_text();self.assertIn('<details open>',html);self.assertIn('a&amp;b.txt',html);self.assertIn('file://',html)
        text=(Path(self.tmp.name)/'structure.txt').read_text(encoding='utf-8-sig');self.assertIn('nested',text);self.assertIn('a&b.txt',text)
    def test_tree_folders_only_export_and_filter(self):
        from tree_export import export_tree
        root=self.fixture_tree();self.w.load(root);settle(self.w);self.w.tree_filter.setText('a&b');self.assertFalse(self.w.tree.topLevelItem(0).isHidden())
        path=Path(self.tmp.name)/'folders.csv';export_tree(self.w.tree_records,root,path,'CSV',False,False);self.assertNotIn('one.txt',path.read_text(encoding='utf-8-sig'));self.assertIn('empty',path.read_text(encoding='utf-8-sig'))
    def test_tree_export_failure_preserves_old_file(self):
        from tree_export import export_tree
        root=self.fixture_tree();self.w.load(root);settle(self.w);path=Path(self.tmp.name)/'existing.csv';path.write_text('ORIGINAL')
        with patch('tree_export.os.replace',side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):export_tree(self.w.tree_records,root,path,'CSV')
        self.assertEqual(path.read_text(),'ORIGINAL');self.assertFalse(list(Path(self.tmp.name).glob('.fwn_export_*')))
    def test_new_folder_invalidates_old_hash(self):
        root=self.fixture_tree();self.w.load(root);settle(self.w);self.w.hash_scan();settle(self.w);other=Path(self.tmp.name)/'other';other.mkdir();self.w.load(other);settle(self.w);self.assertFalse(self.w.hash_records);self.w.hash_scan_btn.click();self.assertIn('0件',self.w.hash_count.text())
    def test_hash_search_and_duplicate_filter(self):
        root=self.fixture_tree();(root/'copy.txt').write_text('ONE');self.w.load(root);settle(self.w);self.w.hash_scan();settle(self.w);self.w.hash_mode.setCurrentText('重複のみ');settle(self.w);self.assertEqual(self.w.hash_list.count(),2);self.w.hash_search.setText('copy.txt');settle(self.w);self.assertEqual(self.w.hash_list.count(),1)
    def test_hash_live_batches_and_large_result_render(self):
        root=Path(self.tmp.name)/'many';root.mkdir()
        for i in range(320):(root/f'{i}.txt').write_text(str(i))
        self.w.load(root);settle(self.w);self.w.hash_scan();worker=self.w.hash_worker;received=[];worker.result_batch.connect(lambda rows:received.append(len(rows)));settle(self.w);self.assertTrue(received);self.assertEqual(sum(received),320);self.assertEqual(self.w.hash_list.count(),320)
    def test_text_preview_and_file_change_guard(self):
        root=self.fixture_tree();self.w.preview(root/'one.txt');settle(self.w);self.assertEqual(self.w.text.toPlainText(),'ONE');self.assertTrue(self.w.edit_btn.isEnabled());self.w.task_active=True
        with patch.object(QMessageBox,'information'),patch.object(self.w,'open_editor') as editor:self.w.edit_current();editor.assert_not_called()
        self.w.task_active=False
    def menu_action(self,widget,text):
        from PySide6.QtCore import QPoint
        from PySide6.QtWidgets import QMenu
        class TestMenu(QMenu):
            def exec(menu,*_):return next(action for action in menu.actions() if action.text()==text)
        with patch('app_main.QMenu',TestMenu),patch.object(QMessageBox,'information'):self.w.file_menu(widget,QPoint(1,1));settle(self.w)
    def test_context_copy_and_move_refresh(self):
        from PySide6.QtWidgets import QListWidgetItem
        root=self.fixture_tree();self.w.load(root);settle(self.w);destination=Path(self.tmp.name)/'destination';destination.mkdir();item=QListWidgetItem('one');item.setData(Qt.UserRole,str(root/'one.txt'));self.w.left.addItem(item);item.setSelected(True)
        with patch.object(QFileDialog,'getExistingDirectory',return_value=str(destination)):self.menu_action(self.w.left,'コピー先を選択')
        self.assertEqual((destination/'one.txt').read_text(),'ONE');self.assertTrue((root/'one.txt').exists())
        with patch.object(QFileDialog,'getExistingDirectory',return_value=str(destination)):self.menu_action(self.w.left,'移動先を選択')
        self.assertFalse((root/'one.txt').exists());self.assertEqual((destination/'one_1.txt').read_text(),'ONE');self.assertEqual(len(self.w.files),1)
    def test_context_duplicate_and_trash(self):
        from PySide6.QtWidgets import QListWidgetItem
        root=self.fixture_tree();self.w.load(root);settle(self.w);item=QListWidgetItem('one');item.setData(Qt.UserRole,str(root/'one.txt'));self.w.left.addItem(item);item.setSelected(True);self.menu_action(self.w.left,'同じ場所に複製');self.assertTrue((root/'one_1.txt').exists())
        with patch('app_main.send2trash',side_effect=lambda path:Path(path).unlink()) as trash,patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):self.menu_action(self.w.left,'ごみ箱へ移す');trash.assert_called_once()
        self.assertFalse((root/'one.txt').exists());self.assertTrue((root/'one_1.txt').exists())
    def test_rename_invalid_regex_and_select_only(self):
        root=self.fixture_tree();self.w.load(root);settle(self.w);self.w.rename_active='regex';self.w.rename_widgets['regex'].setText('[');self.w.rename_preview();self.assertTrue(all(r['status'].startswith('無効') for r in self.w.rename_rows))
        self.w.rename_widgets['regex'].setText('');self.w.rename_selected.setChecked(True);iterator=QTreeWidgetItemIterator(self.w.tree)
        while iterator.value():
            item=iterator.value()
            if item.text(0)=='one.txt':item.setSelected(True)
            iterator+=1
        self.w.rename_active='affix';self.w.rename_widgets['prefix'].setText('new_');self.w.rename_preview();self.assertEqual(len(self.w.rename_rows),1);self.assertEqual(self.w.rename_rows[0]['new'],str(root/'new_one.txt'))

    def test_detached_tree_drop_keeps_tree_context(self):
        root=self.fixture_tree();self.w.popout('tree','構成・ツリー');self.w.show_page('hash');self.drop_on(self.w.tree_root_edit,[root]);self.assertEqual(len(self.w.files),2);self.assertFalse(self.w.hash_records);self.w.floaters['tree'].back()
    def test_tree_excel_formula_name_is_text(self):
        from tree_export import scan_tree,export_tree
        from openpyxl import load_workbook
        root=self.fixture_tree();(root/'=2+3.txt').write_text('TEXT');path=Path(self.tmp.name)/'names.xlsx';export_tree(scan_tree(root),root,path,'Excel');wb=load_workbook(path);cell=next(row[2] for row in wb.active.iter_rows() if row[2].value=='=2+3.txt');self.assertEqual(cell.data_type,'s');wb.close()
    def test_tree_symlink_loop_not_followed(self):
        from tree_export import scan_tree
        root=self.fixture_tree();link=root/'nested'/'loop'
        make_test_symlink(self,link,root,target_is_directory=True)
        records=scan_tree(root);self.assertEqual(len(records),6);self.assertEqual(next(r for r in records if r['name']=='loop')['kind'],'リンク')
    def test_japanese_paths_and_export(self):
        root=Path(self.tmp.name)/'写真と資料';root.mkdir();(root/'設定_日本語.txt').write_text('日本語',encoding='utf-8');self.w.load(root);settle(self.w);self.w.hash_scan_btn.click();settle(self.w);self.assertIn('設定_日本語',self.w.hash_list.item(0).text())
        from tree_export import export_tree
        path=Path(self.tmp.name)/'構成.html';export_tree(self.w.tree_records,root,path,'HTML');self.assertIn('設定_日本語',path.read_text())
    def test_popout_restore(self):
        page=self.w.pages['hash'];self.w.popout('hash','HASH');self.w.floaters['hash'].back();self.assertGreaterEqual(self.w.stack.indexOf(page),0)

if __name__=='__main__':unittest.main(verbosity=2)
